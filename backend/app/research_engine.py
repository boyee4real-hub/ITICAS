
from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timezone, timedelta
from io import BytesIO, StringIO
from math import atan2, cos, radians, sin, sqrt
from statistics import mean, median
from zipfile import ZipFile, ZIP_DEFLATED
from zoneinfo import ZoneInfo
import csv
import json

from openpyxl import Workbook
from PIL import Image, ImageDraw

from sqlalchemy import select, func

from .analytics import analyze_location
from .database import (
    SessionLocal, MonitoringLocation, TrafficObservation, ResearchSurveySession,
    GNSSTrackPoint, CorridorProbe, CorridorProbeObservation
)
from .reporting import export_bytes, _latest_segment_coords, _report_font
from .research_quality import gnss_quality
from .corridor_intelligence import corridor_analysis


def _haversine_m(lat1, lon1, lat2, lon2):
    r=6371000.0
    p1,p2=radians(lat1),radians(lat2)
    dp=radians(lat2-lat1); dl=radians(lon2-lon1)
    a=sin(dp/2)**2+cos(p1)*cos(p2)*sin(dl/2)**2
    return 2*r*atan2(sqrt(a),sqrt(max(0,1-a)))


def _polyline_chainage(coords):
    out=[]; total=0.0
    for i,(lon,lat) in enumerate(coords):
        if i:
            total += _haversine_m(coords[i-1][1],coords[i-1][0],lat,lon)
        out.append({"sequence":i+1,"chainage_m":round(total,2),"longitude":lon,"latitude":lat})
    return out


def _sample_polyline(coords, spacing_m=100):
    if not coords: return []
    chain=_polyline_chainage(coords)
    total=chain[-1]["chainage_m"] if chain else 0
    if total <= 0: return chain[:1]
    targets=[0.0]
    x=float(spacing_m)
    while x < total:
        targets.append(x); x += spacing_m
    if total not in targets: targets.append(total)
    result=[]
    seg_i=1
    for target in targets:
        while seg_i < len(chain) and chain[seg_i]["chainage_m"] < target:
            seg_i += 1
        if seg_i >= len(chain):
            p=chain[-1]
            result.append({"sequence":len(result)+1,"chainage_m":round(target,2),"longitude":p["longitude"],"latitude":p["latitude"]})
            continue
        a=chain[max(0,seg_i-1)]; b=chain[seg_i]
        span=b["chainage_m"]-a["chainage_m"]
        f=0 if span<=0 else (target-a["chainage_m"])/span
        lon=a["longitude"]+f*(b["longitude"]-a["longitude"])
        lat=a["latitude"]+f*(b["latitude"]-a["latitude"])
        result.append({"sequence":len(result)+1,"chainage_m":round(target,2),"longitude":round(lon,7),"latitude":round(lat,7)})
    return result


def ensure_corridor_probes(location_id:int, spacing_m:int=100):
    spacing_m=max(25,min(1000,int(spacing_m)))
    coords=_latest_segment_coords(location_id) or []
    if not coords:
        return {"created":0,"spacing_m":spacing_m,"reason":"No provider segment geometry is stored yet for this road."}
    sampled=_sample_polyline(coords,spacing_m)
    with SessionLocal() as db:
        old=list(db.scalars(select(CorridorProbe).where(CorridorProbe.location_id==location_id).order_by(CorridorProbe.sequence_no)).all())
        old_ids=[r.id for r in old]
        referenced=0
        if old_ids:
            referenced=int(db.scalar(select(func.count(CorridorProbeObservation.id)).where(CorridorProbeObservation.probe_id.in_(old_ids))) or 0)
        # Historical probe observations are scientific evidence. Never delete their parent probes.
        # A repeated 100 m prepare action therefore reuses the established probe network rather
        # than violating the FK or silently destroying/re-parenting historical observations.
        if old and referenced:
            existing_spacing=None
            if len(old)>1:
                diffs=[float(old[i].chainage_m)-float(old[i-1].chainage_m) for i in range(1,len(old)) if float(old[i].chainage_m)>float(old[i-1].chainage_m)]
                existing_spacing=round(median(diffs),1) if diffs else None
            return {
                "created":0,"reused":len(old),"spacing_m":existing_spacing or spacing_m,
                "requested_spacing_m":spacing_m,"segment_length_m":float(old[-1].chainage_m) if old else 0,
                "historical_observations_preserved":referenced,
                "reason":f"Existing corridor probes were reused because {referenced} historical probe observation(s) reference them; no research evidence was deleted."
            }
        for r in old: db.delete(r)
        db.flush()
        for p in sampled:
            db.add(CorridorProbe(location_id=location_id,sequence_no=p["sequence"],chainage_m=p["chainage_m"],
                                 latitude=p["latitude"],longitude=p["longitude"],source_geometry="provider_segment"))
        db.commit()
    return {"created":len(sampled),"reused":0,"spacing_m":spacing_m,"segment_length_m":sampled[-1]["chainage_m"] if sampled else 0}


def survey_summary(location_id:int|None=None):
    with SessionLocal() as db:
        q=select(ResearchSurveySession)
        if location_id is not None:
            q=q.where(ResearchSurveySession.location_id==location_id)
        sessions=list(db.scalars(q.order_by(ResearchSurveySession.started_at.desc())).all())
        session_ids=[x.id for x in sessions]
        pts=[]
        if session_ids:
            pts=list(db.scalars(select(GNSSTrackPoint).where(GNSSTrackPoint.session_id.in_(session_ids)).order_by(GNSSTrackPoint.captured_at)).all())
    valid_pts=[]; invalid_pts=[]
    for p in pts:
        sid=next((x for x in sessions if x.id==p.session_id),None)
        q=gnss_quality(sid.location_id,p.latitude,p.longitude,p.accuracy_m) if sid and sid.location_id is not None else {"valid":True}
        (valid_pts if q.get("valid") else invalid_pts).append(p)
    acc=[p.accuracy_m for p in valid_pts if p.accuracy_m is not None]
    sp=[p.speed_mps*3.6 for p in valid_pts if p.speed_mps is not None]
    return {
        "sessions":len(sessions),"points":len(valid_pts),"stored_points_total":len(pts),"excluded_out_of_corridor_points":len(invalid_pts),
        "mean_horizontal_accuracy_m":round(mean(acc),2) if acc else None,
        "median_speed_kmh":round(median(sp),2) if sp else None,
        "first_point_at":valid_pts[0].captured_at.isoformat() if valid_pts else None,
        "last_point_at":valid_pts[-1].captured_at.isoformat() if valid_pts else None,
        "active_session_ids":[x.id for x in sessions if x.status=="active"],
    }


def objective_matrix(location_id:int,days:int=30):
    a=analyze_location(location_id,days)
    if not a: return None
    survey_info=survey_summary(location_id)
    survey_count=survey_info.get("points",0)
    with SessionLocal() as db:
        probe_count=db.scalar(select(func.count(CorridorProbe.id)).where(CorridorProbe.location_id==location_id)) or 0
    seg=_latest_segment_coords(location_id) or []
    coverage=a["coverage"]["combined_hourly_coverage_pct"]
    rows=[
        {
            "objective":"Collect GPS/GNSS coordinates of congestion points",
            "status":"complete" if survey_count>=10 else "partial" if survey_count else "not_started",
            "evidence":f"{survey_count} valid field GNSS points; {survey_info.get('excluded_out_of_corridor_points',0)} out-of-corridor point(s) excluded",
            "next_action":"Collect a field/mobile GNSS track with Research Workbench." if survey_count<10 else "Field coordinate evidence available."
        },
        {
            "objective":"Digitize the road network / selected corridor",
            "status":"complete" if len(seg)>=2 else "partial" if a["location"].get("road_name") else "not_started",
            "evidence":f"{len(seg)} provider road-geometry vertices available",
            "next_action":"Acquire live traffic once to retain provider segment geometry." if len(seg)<2 else "Provider centerline can be exported to GIS formats."
        },
        {
            "objective":"Analyze congestion intensity using GIS/statistical tools",
            "status":"complete" if a["data_quality"]["observations"]>=12 else "partial" if a["data_quality"]["observations"] else "not_started",
            "evidence":f"{a['data_quality']['observations']} live observations; {coverage}% hourly evidence coverage",
            "next_action":"Continue monitoring to improve temporal coverage." if coverage<20 else "Intensity analytics available."
        },
        {
            "objective":"Produce traffic-congestion hotspot maps",
            "status":"complete" if probe_count>=5 and coverage>=20 else "screening_only" if a["data_quality"]["observations"] else "not_started",
            "evidence":f"{probe_count} corridor probes; {coverage}% temporal coverage",
            "next_action":"Create corridor probes and collect repeated point-specific observations before claiming intra-road statistical hotspots." if probe_count<5 or coverage<20 else "Evidence supports stronger hotspot analysis."
        },
    ]
    complete=sum(1 for r in rows if r["status"]=="complete")
    return {"location":a["location"],"period_days":days,"objectives":rows,"complete_objectives":complete,"total_objectives":len(rows),
            "overall_completion_pct":round(100*complete/len(rows),1)}


def research_catalog(location_id:int,days:int=30):
    a=analyze_location(location_id,days)
    if not a:return None
    obj=objective_matrix(location_id,days)
    ss=survey_summary(location_id)
    outputs=[
        ("Research objective matrix","table","Always"),
        ("GNSS survey-session register","table","When survey sessions exist"),
        ("GNSS trajectory coordinates","table + GIS","When GNSS points exist"),
        ("Selected road/provider centerline","GIS map","When segment geometry exists"),
        ("Corridor chainage/probe table","table + GIS","When probes are generated"),
        ("Raw live traffic observations","table","When observations exist"),
        ("True hourly traffic time series","table + chart","Always; gaps remain missing"),
        ("Clock-hour congestion profile","chart","When traffic evidence exists"),
        ("Daily congestion profile","chart","When multi-day evidence exists"),
        ("Weekday/weekend comparison","table + chart","When day-type evidence exists"),
        ("Speed vs free-flow profile","chart","When paired speed data exist"),
        ("Speed-reduction profile","chart","When paired speed data exist"),
        ("Delay distribution/profile","chart","When delay data exist"),
        ("Travel-time reliability TTI/PTI/Buffer Index","table + chart","When travel times exist"),
        ("Congestion-severity distribution","chart","When CI exists"),
        ("Congestion persistence/episode analysis","table + chart","When hourly evidence exists"),
        ("Peak-hour / AM / PM analysis","table + chart","When temporal evidence exists"),
        ("Day × hour congestion heatmap","figure","When hourly evidence exists"),
        ("Evidence completeness/provenance timeline","figure","Always"),
        ("Road decision dashboard","figure","When traffic evidence exists"),
        ("Incident context","table","When incident data exist"),
        ("Weather-traffic relationship","analysis","When synchronized weather exists"),
        ("Single-road congestion evidence map","map","When road geometry/evidence exist"),
        ("Multi-route comparison / intervention ranking","map + table","Optional network mode"),
        ("Gi* / Local Moran / KDE suitability gate","scientific diagnostic","Only when spatial evidence is adequate"),
        ("Research-ready DOCX/PDF report","report","Always"),
        ("CSV/XLSX/JSON reproducibility package","data package","Always"),
        ("GeoJSON/KML/GPX GIS exchange package","GIS package","When geometry exists"),
    ]
    return {
        "location":a["location"],"period_days":days,
        "survey":ss,"objectives":obj,
        "outputs":[{"name":n,"type":t,"availability":av} for n,t,av in outputs],
        "output_count":len(outputs),
        "scientific_rule":"ITICAS generates every valid output supported by the available evidence; it does not fabricate unsupported variables, hours, traffic volumes, directions or hotspot significance."
    }


def _csv_bytes(rows,fields):
    s=StringIO(newline="");w=csv.DictWriter(s,fieldnames=fields);w.writeheader()
    for r in rows:w.writerow({k:r.get(k) for k in fields})
    return s.getvalue().encode("utf-8-sig")


def _geojson_feature_collection(features):
    return json.dumps({"type":"FeatureCollection","features":features},indent=2).encode("utf-8")


def _gnss_rows(location_id:int):
    with SessionLocal() as db:
        sessions=list(db.scalars(select(ResearchSurveySession).where(ResearchSurveySession.location_id==location_id)).all())
        ids=[s.id for s in sessions]
        pts=list(db.scalars(select(GNSSTrackPoint).where(GNSSTrackPoint.session_id.in_(ids)).order_by(GNSSTrackPoint.captured_at)).all()) if ids else []
    rows=[]
    for p in pts:
        q=gnss_quality(location_id,p.latitude,p.longitude,p.accuracy_m)
        rows.append({"session_id":p.session_id,"captured_at":p.captured_at.isoformat(),"latitude":p.latitude,"longitude":p.longitude,
           "altitude_m":p.altitude_m,"accuracy_m":p.accuracy_m,"speed_kmh":round(p.speed_mps*3.6,2) if p.speed_mps is not None else None,
           "heading_deg":p.heading_deg,"traffic_state":p.traffic_state,"queue_length_m":p.queue_length_m,"notes":p.notes,"source":p.source,
           "study_spatial_valid":bool(q.get("valid")),"distance_to_study_corridor_m":q.get("distance_m"),"quality_note":q.get("reason")})
    return rows


def _gnss_exports(location_id:int):
    rows=_gnss_rows(location_id); valid_rows=[r for r in rows if r.get("study_spatial_valid")]
    fields=list(rows[0].keys()) if rows else ["session_id","captured_at","latitude","longitude","altitude_m","accuracy_m","speed_kmh","heading_deg","traffic_state","queue_length_m","notes","source","study_spatial_valid","distance_to_study_corridor_m","quality_note"]
    csvb=_csv_bytes(valid_rows,fields)
    feats=[{"type":"Feature","geometry":{"type":"Point","coordinates":[r["longitude"],r["latitude"]]},"properties":{k:v for k,v in r.items() if k not in ("latitude","longitude")}} for r in valid_rows]
    if valid_rows:
        feats.insert(0,{"type":"Feature","geometry":{"type":"LineString","coordinates":[[r["longitude"],r["latitude"]] for r in valid_rows]},"properties":{"name":"GNSS survey trajectory — spatially valid points only","point_count":len(valid_rows)}})
    gpx=['<?xml version="1.0" encoding="UTF-8"?><gpx version="1.1" creator="ITICAS" xmlns="http://www.topografix.com/GPX/1/1"><trk><name>ITICAS GNSS Traffic Survey - Valid Study Points</name><trkseg>']
    for r in valid_rows:gpx.append(f'<trkpt lat="{r["latitude"]}" lon="{r["longitude"]}"><time>{r["captured_at"]}</time></trkpt>')
    gpx.append('</trkseg></trk></gpx>')
    return csvb,_geojson_feature_collection(feats),"".join(gpx).encode("utf-8")


def _gnss_quarantine_export(location_id:int):
    rows=[r for r in _gnss_rows(location_id) if not r.get("study_spatial_valid")]
    fields=list(rows[0].keys()) if rows else ["session_id","captured_at","latitude","longitude","altitude_m","accuracy_m","speed_kmh","heading_deg","traffic_state","queue_length_m","notes","source","study_spatial_valid","distance_to_study_corridor_m","quality_note"]
    return _csv_bytes(rows,fields),len(rows)

def _road_exports(location_id:int):
    coords=_latest_segment_coords(location_id) or []
    feature={"type":"Feature","geometry":{"type":"LineString","coordinates":[[x,y] for x,y in coords]},"properties":{"location_id":location_id,"source":"TomTom Traffic API provider segment geometry"}}
    geo=_geojson_feature_collection([feature] if coords else [])
    kml=['<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document><Placemark><name>ITICAS selected road segment</name><LineString><coordinates>']
    kml.append(" ".join(f"{x},{y},0" for x,y in coords))
    kml.append('</coordinates></LineString></Placemark></Document></kml>')
    chain=_polyline_chainage(coords)
    chain_csv=_csv_bytes(chain,["sequence","chainage_m","longitude","latitude"])
    return geo,"".join(kml).encode("utf-8"),chain_csv


def _simple_figure(title,subtitle,rows,value_key,label_key,unit=""):
    W,H=1800,1050;img=Image.new("RGB",(W,H),"white");d=ImageDraw.Draw(img)
    d.text((80,35),title,font=_report_font(42,True),fill=(12,20,32));d.text((80,95),subtitle,font=_report_font(22),fill=(70,80,95))
    left,top,right,bottom=150,180,1720,850;d.rectangle((left,top,right,bottom),outline=(90,105,120),width=2)
    vals=[r.get(value_key) for r in rows if r.get(value_key) is not None]
    if not vals:d.text((left+40,top+60),"Insufficient evidence for this figure.",font=_report_font(28,True),fill=(120,60,60))
    else:
        vmax=max(vals) or 1; shown=rows[-30:] if len(rows)>30 else rows; bw=max(8,(right-left-40)//max(1,len(shown)))
        tick_every=max(1,len(shown)//10)
        for i,r in enumerate(shown):
            v=r.get(value_key); x=left+20+i*bw
            if v is not None:
                bh=(bottom-top-100)*(float(v)/vmax); y=bottom-50-bh
                ci=float(v) if 'congestion' in value_key else None
                col=(64,180,140) if ci is None or ci<.2 else (245,205,70) if ci<.4 else (245,135,55) if ci<.65 else (205,50,60)
                d.rectangle((x,y,x+bw-3,bottom-50),fill=col)
            if i%tick_every==0 or i==len(shown)-1:
                lab=str(r.get(label_key,'')); lab=lab[11:16] if 'T' in lab else lab[-10:]
                d.text((x,bottom-35),lab,font=_report_font(13),fill=(55,65,75))
        d.text((left+10,bottom+28),f"X-axis: {label_key.replace('_',' ')} / interval labels · Maximum displayed: {round(vmax,3)}{unit}",font=_report_font(17),fill=(60,70,80))
    d.text((80,975),"Developed by Dr. Oyeyode A.O. MNIS, MGEOSON, MNAG",font=_report_font(18,True),fill=(20,90,70))
    bio=BytesIO();img.save(bio,"PNG",dpi=(300,300));return bio.getvalue()

def _corridor_congestion_map(location_id:int,days:int):
    ca=corridor_analysis(location_id,days); rows=[r for r in ca.get("rows",[]) if r.get("latitude") is not None and r.get("longitude") is not None]
    W,H=1800,1100;img=Image.new("RGB",(W,H),"white");d=ImageDraw.Draw(img)
    d.text((80,35),"CORRIDOR CONGESTION SEVERITY MAP",font=_report_font(42,True),fill=(12,20,32));d.text((80,95),f"{days}-day probe evidence · each route section is coloured by mean congestion",font=_report_font(22),fill=(70,80,95))
    if len(rows)<2:
        d.text((120,220),"Insufficient repeated corridor-probe evidence for a route-severity map.",font=_report_font(28,True),fill=(130,60,60))
    else:
        lons=[r['longitude'] for r in rows]; lats=[r['latitude'] for r in rows]; minx,maxx=min(lons),max(lons); miny,maxy=min(lats),max(lats)
        pad=.08; dx=max(maxx-minx,1e-6);dy=max(maxy-miny,1e-6); minx-=dx*pad;maxx+=dx*pad;miny-=dy*pad;maxy+=dy*pad
        def xy(r):return (150+(r['longitude']-minx)/(maxx-minx)*1450, 900-(r['latitude']-miny)/(maxy-miny)*700)
        def col(ci):
            if ci is None:return (150,160,175)
            if ci<.20:return (65,200,170)
            if ci<.35:return (250,210,55)
            if ci<.50:return (245,150,55)
            if ci<.70:return (230,70,60)
            return (150,35,55)
        for a,b in zip(rows,rows[1:]):
            ci=[x for x in (a.get('mean_congestion_index'),b.get('mean_congestion_index')) if x is not None]; c=col(mean(ci) if ci else None);d.line((*xy(a),*xy(b)),fill=c,width=18)
        for i,r in enumerate(rows):
            x,y=xy(r);c=col(r.get('mean_congestion_index'));d.ellipse((x-7,y-7,x+7,y+7),fill=c,outline='white',width=2)
            if i in (0,len(rows)-1) or r.get('screening_hotspot'):
                d.text((x+10,y-15),f"P{r['sequence_no']} · {r['chainage_m']:.0f}m · CI {r.get('mean_congestion_index') if r.get('mean_congestion_index') is not None else '—'}",font=_report_font(15,True),fill=(30,40,50))
        # Most congested evidence location
        valid=[r for r in rows if r.get('mean_congestion_index') is not None]
        if valid:
            worst=max(valid,key=lambda r:r['mean_congestion_index']); x,y=xy(worst);d.ellipse((x-22,y-22,x+22,y+22),outline=(120,0,35),width=6);d.text((110,960),f"Most congested observed probe: P{worst['sequence_no']} at chainage {worst['chainage_m']:.0f} m · mean CI {worst['mean_congestion_index']} · {worst['severity']}",font=_report_font(20,True),fill=(90,20,35))
        lx=110; ly=1010
        for lab,c in [('Free',(65,200,170)),('Light',(250,210,55)),('Moderate',(245,150,55)),('Heavy',(230,70,60)),('Severe',(150,35,55)),('No evidence',(150,160,175))]:
            d.rectangle((lx,ly,lx+20,ly+20),fill=c);d.text((lx+28,ly-2),lab,font=_report_font(15),fill=(30,40,50));lx+=220
    bio=BytesIO();img.save(bio,"PNG",dpi=(300,300));return bio.getvalue()



ANALYSIS_TZ=ZoneInfo("Africa/Lagos")

def _severity_color(ci):
    if ci is None:return (150,160,175)
    if ci<.20:return (45,212,191)
    if ci<.35:return (250,204,21)
    if ci<.50:return (251,146,60)
    if ci<.70:return (239,68,68)
    return (185,28,28)

def _daily_hourly_timeline_figure(day, location_name):
    rows=list(day.get("hourly") or [])[:24]
    W,H=1800,1000; img=Image.new("RGB",(W,H),"white"); d=ImageDraw.Draw(img)
    d.text((80,35),f"DAILY HOURLY CONGESTION — {location_name}",font=_report_font(40,True),fill=(12,20,32))
    d.text((80,92),f"{day.get('date')} · fixed 24-hour West Africa Time axis · gaps mean no evidence",font=_report_font(22),fill=(70,80,95))
    L,T,R,B=150,180,1720,820; PW=R-L; PH=B-T
    for i in range(5):
        yy=T+PH*i/4; val=1-i/4; d.line((L,yy,R,yy),fill=(220,226,234),width=2); d.text((80,yy-10),f"{val:.2f}",font=_report_font(16),fill=(70,80,95))
    d.line((L,T,L,B),fill=(60,70,80),width=3);d.line((L,B,R,B),fill=(60,70,80),width=3)
    pts=[]
    for h in range(24):
        x=L+(PW*h/23); r=rows[h] if h<len(rows) else None; ci=(r or {}).get("mean_congestion_index")
        if ci is not None:
            y=B-max(0,min(1,float(ci)))*PH; pts.append((x,y)); d.ellipse((x-7,y-7,x+7,y+7),fill=_severity_color(float(ci)),outline="white",width=2)
        else:
            if len(pts)>1:d.line(pts,fill=(30,150,220),width=5)
            pts=[]
        if h%2==0:d.text((x-22,B+18),f"{h:02d}:00",font=_report_font(14),fill=(55,65,75))
    if len(pts)>1:d.line(pts,fill=(30,150,220),width=5)
    peak=day.get("peak_hour")
    if peak:
        try:
            ph=datetime.fromisoformat(str(peak)).astimezone(ANALYSIS_TZ).hour; px=L+(PW*ph/23); d.line((px,T,px,B),fill=(220,25,65),width=3); d.text((min(px+8,R-220),T+10),f"Peak: {ph:02d}:00 WAT · CI {day.get('peak_congestion_index')}",font=_report_font(17,True),fill=(160,20,50))
        except Exception: pass
    d.text((620,890),"Hour of day — West Africa Time (00:00–23:00)",font=_report_font(19,True),fill=(45,55,70))
    d.text((80,950),"Daily result is produced before the composite multi-day analysis.",font=_report_font(17),fill=(60,70,80))
    bio=BytesIO();img.save(bio,"PNG",dpi=(300,300));return bio.getvalue()

def _probe_hourly_stats(location_id:int):
    with SessionLocal() as db:
        probes={p.id:p for p in db.scalars(select(CorridorProbe).where(CorridorProbe.location_id==location_id)).all()}
        obs=list(db.scalars(select(CorridorProbeObservation).where(CorridorProbeObservation.location_id==location_id).order_by(CorridorProbeObservation.observed_at)).all())
    groups=defaultdict(lambda:defaultdict(list))
    for o in obs:
        if o.congestion_index is None: continue
        dt=o.observed_at.replace(tzinfo=timezone.utc) if o.observed_at.tzinfo is None else o.observed_at
        local=dt.astimezone(ANALYSIS_TZ); groups[(local.date().isoformat(),local.hour)][o.probe_id].append(float(o.congestion_index))
    return probes,groups

def _probe_map_image(probes, stats, title, subtitle):
    rows=[]
    for pid,p in probes.items():
        if p.latitude is None or p.longitude is None: continue
        vals=stats.get(pid) or []; rows.append({"p":p,"ci":mean(vals) if vals else None})
    rows.sort(key=lambda x:x["p"].sequence_no)
    W,H=1800,1050;img=Image.new("RGB",(W,H),"white");d=ImageDraw.Draw(img)
    d.text((80,35),title,font=_report_font(38,True),fill=(12,20,32));d.text((80,92),subtitle,font=_report_font(21),fill=(70,80,95))
    valid=[r for r in rows if r["ci"] is not None]
    if len(rows)<2:
        d.text((120,220),"No usable corridor geometry for this map.",font=_report_font(28,True),fill=(130,60,60))
    else:
        lons=[r['p'].longitude for r in rows];lats=[r['p'].latitude for r in rows];minx,maxx=min(lons),max(lons);miny,maxy=min(lats),max(lats);dx=max(maxx-minx,1e-6);dy=max(maxy-miny,1e-6);minx-=dx*.08;maxx+=dx*.08;miny-=dy*.08;maxy+=dy*.08
        def xy(r):return (150+(r['p'].longitude-minx)/(maxx-minx)*1450,880-(r['p'].latitude-miny)/(maxy-miny)*680)
        for a,b in zip(rows,rows[1:]):
            vals=[x for x in (a['ci'],b['ci']) if x is not None]; d.line((*xy(a),*xy(b)),fill=_severity_color(mean(vals) if vals else None),width=18)
        for r in rows:
            x,y=xy(r);d.ellipse((x-7,y-7,x+7,y+7),fill=_severity_color(r['ci']),outline='white',width=2)
        if valid:
            w=max(valid,key=lambda r:r['ci']);x,y=xy(w);d.ellipse((x-22,y-22,x+22,y+22),outline=(120,0,35),width=6);d.text((110,930),f"Strongest observed location: probe {w['p'].sequence_no} · chainage {w['p'].chainage_m:.0f} m · CI {w['ci']:.3f}",font=_report_font(20,True),fill=(90,20,35))
        else:d.text((110,930),"No probe-specific congestion evidence exists for this interval; route is shown as no evidence.",font=_report_font(19),fill=(80,90,100))
    lx=100;ly=985
    for lab,c in [('Free',(45,212,191)),('Light',(250,204,21)),('Moderate',(251,146,60)),('Heavy',(239,68,68)),('Severe',(185,28,28)),('No evidence',(150,160,175))]:d.rectangle((lx,ly,lx+20,ly+20),fill=c);d.text((lx+28,ly-2),lab,font=_report_font(14),fill=(30,40,50));lx+=215
    bio=BytesIO();img.save(bio,"PNG",dpi=(300,300));return bio.getvalue()

def _daily_probe_products(location_id:int, date_str:str):
    probes,groups=_probe_hourly_stats(location_id); hourly={h:groups.get((date_str,h),{}) for h in range(24)}; daily=defaultdict(list)
    for stats in hourly.values():
        for pid,vals in stats.items(): daily[pid].extend(vals)
    products={"daily_map":_probe_map_image(probes,daily,f"DAILY CORRIDOR CONGESTION MAP — {date_str}","Probe-specific congestion by route location; daily evidence only"),"hourly_maps":{}}
    for h,stats in hourly.items():
        if any(stats.values()): products["hourly_maps"][h]=_probe_map_image(probes,stats,f"HOURLY CORRIDOR CONGESTION MAP — {date_str}",f"{h:02d}:00–{(h+1)%24:02d}:00 WAT · probe-specific evidence only")
    return products

def _corridor_observation_export(location_id:int):
    with SessionLocal() as db:
        probes={p.id:p for p in db.scalars(select(CorridorProbe).where(CorridorProbe.location_id==location_id)).all()}
        obs=list(db.scalars(select(CorridorProbeObservation).where(CorridorProbeObservation.location_id==location_id).order_by(CorridorProbeObservation.observed_at)).all())
    rows=[]
    for o in obs:
        p=probes.get(o.probe_id)
        rows.append({"probe_id":o.probe_id,"sequence_no":p.sequence_no if p else None,"chainage_m":p.chainage_m if p else None,
                     "observed_at":o.observed_at.isoformat(),"current_speed_kmh":o.current_speed_kmh,
                     "free_flow_speed_kmh":o.free_flow_speed_kmh,"congestion_index":o.congestion_index,
                     "delay_seconds":o.delay_seconds,"confidence":o.confidence,"provider":o.provider,
                     "provider_status":o.provider_status,"error_category":o.error_category})
    fields=list(rows[0].keys()) if rows else ["probe_id","sequence_no","chainage_m","observed_at","current_speed_kmh","free_flow_speed_kmh","congestion_index","delay_seconds","confidence","provider","provider_status","error_category"]
    return _csv_bytes(rows,fields)

def build_research_workbook(location_id:int,days:int):
    a=analyze_location(location_id,days); obj=objective_matrix(location_id,days); cat=research_catalog(location_id,days)
    wb=Workbook();ws=wb.active;ws.title="Research Objectives"
    ws.append(["Objective","Status","Evidence","Next Action"])
    for r in obj["objectives"]:ws.append([r["objective"],r["status"],r["evidence"],r["next_action"]])
    ws2=wb.create_sheet("Traffic Summary");ws2.append(["Metric","Value"])
    for k,v in a["summary"].items():ws2.append([k,json.dumps(v) if isinstance(v,(dict,list)) else v])
    ws3=wb.create_sheet("Hourly Evidence");ws3.append(list(a["hourly_timeseries"][0].keys()) if a["hourly_timeseries"] else ["hour_start"])
    for r in a["hourly_timeseries"]:ws3.append(list(r.values()))
    ws4=wb.create_sheet("Day by Day");ws4.append(["Date","Hours With Evidence","Coverage %","Mean CI","Max CI","Mean Speed km/h","Mean Delay s","Peak Time","Peak CI","Peak Category"])
    for r in a.get("day_by_day",[]):
        ws4.append([r.get("date"),r.get("hours_with_evidence"),r.get("coverage_pct"),r.get("mean_congestion_index"),r.get("max_congestion_index"),r.get("mean_speed_kmh"),r.get("mean_delay_seconds"),r.get("peak_hour"),r.get("peak_congestion_index"),r.get("peak_category")])
        ds=wb.create_sheet(("D "+str(r.get("date")))[:31]);ds.append(["Hour WAT","Status","Samples","Mean Speed km/h","Free-flow km/h","CI","Delay s","Source"])
        for h,x in enumerate(r.get("hourly") or []):ds.append([f"{h:02d}:00-{(h+1)%24:02d}:00",x.get("data_status"),x.get("samples"),x.get("mean_speed_kmh"),x.get("mean_free_flow_speed_kmh"),x.get("mean_congestion_index"),x.get("mean_delay_seconds"),x.get("source")])
    ws5=wb.create_sheet("Output Catalogue");ws5.append(["Output","Type","Availability"])
    for r in cat["outputs"]:ws5.append([r["name"],r["type"],r["availability"]])
    bio=BytesIO();wb.save(bio);return bio.getvalue()


def build_research_package(location_id:int,days:int=30):
    a=analyze_location(location_id,days)
    if not a: raise ValueError("Location not found")
    obj=objective_matrix(location_id,days);cat=research_catalog(location_id,days)
    gnss_csv,gnss_geo,gpx=_gnss_exports(location_id)
    gnss_quarantine,gnss_quarantine_count=_gnss_quarantine_export(location_id)
    road_geo,road_kml,chain_csv=_road_exports(location_id)
    safe="".join(c if c.isalnum() else "_" for c in a["location"]["name"]).strip("_") or f"location_{location_id}"
    bio=BytesIO()
    with ZipFile(bio,"w",ZIP_DEFLATED) as z:
        z.writestr(f"{safe}_research_catalog.json",json.dumps(cat,indent=2,default=str))
        z.writestr(f"{safe}_objective_matrix.csv",_csv_bytes(obj["objectives"],["objective","status","evidence","next_action"]))
        z.writestr(f"{safe}_research_workbook.xlsx",build_research_workbook(location_id,days))
        z.writestr(f"{safe}_gnss_track_points_VALID_ONLY.csv",gnss_csv)
        z.writestr(f"{safe}_gnss_rejected_or_off_corridor_QUARANTINE.csv",gnss_quarantine)
        z.writestr(f"{safe}_gnss_trajectory.geojson",gnss_geo)
        z.writestr(f"{safe}_gnss_trajectory.gpx",gpx)
        z.writestr(f"{safe}_road_segment.geojson",road_geo)
        z.writestr(f"{safe}_road_segment.kml",road_kml)
        z.writestr(f"{safe}_road_chainage.csv",chain_csv)
        z.writestr(f"{safe}_corridor_probe_observations.csv",_corridor_observation_export(location_id))
        z.writestr(f"{safe}_hourly_timeseries.csv",_csv_bytes(a["hourly_timeseries"],list(a["hourly_timeseries"][0].keys()) if a["hourly_timeseries"] else ["hour_start"]))
        z.writestr(f"{safe}_daily_profile.csv",_csv_bytes(a["daily_profile"],["date","mean_congestion_index","samples"]))
        daily_summary=[{k:v for k,v in r.items() if k!="hourly"} for r in a.get("day_by_day",[])]
        z.writestr(f"{safe}_day_by_day_analysis.csv",_csv_bytes(daily_summary,list(daily_summary[0].keys()) if daily_summary else ["date"]))
        for r in a.get("day_by_day",[]):
            daydir=f"01_DAILY_ANALYSIS/{r['date']}"
            z.writestr(f"{daydir}/{safe}_{r['date']}_hourly_00_to_23_WAT.csv",_csv_bytes(r.get("hourly",[]),list(r.get("hourly",[{}])[0].keys()) if r.get("hourly") else ["hour_start"]))
            z.writestr(f"{daydir}/{safe}_{r['date']}_hourly_congestion_timeline.png",_daily_hourly_timeline_figure(r,a["location"]["name"]))
            pp=_daily_probe_products(location_id,r['date']); z.writestr(f"{daydir}/{safe}_{r['date']}_corridor_congestion_map.png",pp["daily_map"])
            for h,img in pp["hourly_maps"].items():z.writestr(f"{daydir}/hourly_corridor_maps/{h:02d}00_{(h+1)%24:02d}00_WAT.png",img)
        z.writestr(f"{safe}_weekday_profile.csv",_csv_bytes(a["weekday_profile"],["day","mean_congestion_index","samples"]))
        z.writestr(f"{safe}_raw_observations.csv",_csv_bytes(a["raw_observations"],list(a["raw_observations"][0].keys()) if a["raw_observations"] else ["id"]))
        # Reuse publication report outputs.
        for fmt in ("docx","pdf","csv","xlsx","json","svg","png"):
            data,mime,name=export_bytes(location_id,days,fmt)
            z.writestr(name,data)
        # Reuse decision atlas products individually.
        maps_blob,_,_=export_bytes(location_id,days,"maps")
        with ZipFile(BytesIO(maps_blob)) as mz:
            for n in mz.namelist():
                z.writestr(f"decision_atlas/{n}",mz.read(n))
        # Additional research figures.
        z.writestr(f"02_COMPOSITE_ANALYSIS/research_figures/{safe}_corridor_congestion_severity_map.png",_corridor_congestion_map(location_id,days))
        z.writestr(f"research_figures/{safe}_corridor_congestion_severity_map.png",_corridor_congestion_map(location_id,days))
        z.writestr(f"02_COMPOSITE_ANALYSIS/research_figures/{safe}_daily_congestion_profile.png",
                   _simple_figure(f"DAILY CONGESTION PROFILE — {a['location']['name']}",f"Last {days} days; actual evidence only",a["daily_profile"],"mean_congestion_index","date"))
        z.writestr(f"research_figures/{safe}_daily_congestion_profile.png",
                   _simple_figure(f"DAILY CONGESTION PROFILE — {a['location']['name']}",f"Last {days} days; actual evidence only",a["daily_profile"],"mean_congestion_index","date"))
        z.writestr(f"02_COMPOSITE_ANALYSIS/research_figures/{safe}_clock_hour_profile.png",
                   _simple_figure(f"CLOCK-HOUR CONGESTION PROFILE — {a['location']['name']}","Mean congestion index by clock hour",a["hourly_profile"],"mean_congestion_index","hour"))
        z.writestr(f"research_figures/{safe}_clock_hour_profile.png",
                   _simple_figure(f"CLOCK-HOUR CONGESTION PROFILE — {a['location']['name']}","Mean congestion index by clock hour",a["hourly_profile"],"mean_congestion_index","hour"))
        evidence=[{"label":r["hour_start"],"value":1 if r["data_status"]!="missing" else 0} for r in a["hourly_timeseries"]]
        z.writestr(f"02_COMPOSITE_ANALYSIS/research_figures/{safe}_evidence_availability.png",
                   _simple_figure(f"EVIDENCE AVAILABILITY — {a['location']['name']}","1 = observed/historical evidence; 0 = missing",evidence,"value","label"))
        z.writestr(f"research_figures/{safe}_evidence_availability.png",
                   _simple_figure(f"EVIDENCE AVAILABILITY — {a['location']['name']}","1 = observed/historical evidence; 0 = missing",evidence,"value","label"))
        manifest={
            "application":"Nigeria Traffic Intelligence and Congestion Analysis System",
            "research_direction":"GNSS + GIS traffic-congestion research workflow",
            "location":a["location"],"period_days":days,
            "objective_completion":obj,
            "scientific_integrity":"Missing evidence is not fabricated. Daily hourly analysis/maps are produced before composite results. Browser GNSS outside Nigeria/off-corridor is quarantined and excluded from valid study evidence.",
            "gnss_quarantined_points":gnss_quarantine_count,
            "developer":"Developed by Dr. Oyeyode A.O. MNIS, MGEOSON, MNAG",
        }
        z.writestr("README_RESEARCH_PACKAGE.json",json.dumps(manifest,indent=2))
    return f"ITICAS_{safe}_{days}d_GNSS_GIS_RESEARCH_PACKAGE.zip",bio.getvalue(),"application/zip"


def build_individual_tool_output(location_id:int, days:int, tool:str):
    """Return a dedicated machine-readable/downloadable output for one research tool."""
    tool=(tool or "").strip().lower()
    if tool=="road":
        geo,kml,chain=_road_exports(location_id)
        bio=BytesIO()
        with ZipFile(bio,"w",ZIP_DEFLATED) as z:
            z.writestr("road_segment.geojson",geo)
            z.writestr("road_segment.kml",kml)
            z.writestr("road_chainage.csv",chain)
            with SessionLocal() as db:
                probes=list(db.scalars(select(CorridorProbe).where(CorridorProbe.location_id==location_id).order_by(CorridorProbe.sequence_no)).all())
            rows=[{"sequence_no":p.sequence_no,"chainage_m":p.chainage_m,"latitude":p.latitude,"longitude":p.longitude,"source_geometry":p.source_geometry} for p in probes]
            z.writestr("corridor_probes.csv",_csv_bytes(rows,["sequence_no","chainage_m","latitude","longitude","source_geometry"]))
        return "ITICAS_road_corridor_output.zip",bio.getvalue(),"application/zip"
    if tool=="scan":
        data=_corridor_observation_export(location_id)
        return "ITICAS_corridor_scan_observations.csv",data,"text/csv"
    if tool=="gnss":
        csvb,geo,gpx=_gnss_exports(location_id); quarantine,qcount=_gnss_quarantine_export(location_id)
        bio=BytesIO()
        with ZipFile(bio,"w",ZIP_DEFLATED) as z:
            z.writestr("gnss_track_points_VALID_ONLY.csv",csvb)
            z.writestr("gnss_rejected_or_off_corridor_QUARANTINE.csv",quarantine)
            z.writestr("gnss_trajectory_VALID_ONLY.geojson",geo)
            z.writestr("gnss_trajectory_VALID_ONLY.gpx",gpx)
            z.writestr("README_GNSS_QUALITY.txt",f"Primary GNSS outputs contain only spatially valid study-corridor points. {qcount} stored point(s) were quarantined as outside Nigeria/off-corridor and retained only for provenance. Browser GNSS reports the physical device position; ITICAS never substitutes the selected road coordinate.\n".encode("utf-8"))
        return "ITICAS_GNSS_field_survey_output.zip",bio.getvalue(),"application/zip"
    if tool=="objectives":
        obj=objective_matrix(location_id,days)
        if not obj: raise ValueError("Location not found")
        return "ITICAS_research_objective_matrix.csv",_csv_bytes(obj["objectives"],["objective","status","evidence","next_action"]),"text/csv"
    raise ValueError("Unknown research tool output.")
